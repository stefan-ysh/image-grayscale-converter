import pandas as pd
import numpy as np
import wx
from datetime import datetime
from tkinter import filedialog, messagebox
from openpyxl.chart import LineChart, Reference
from utils.show_progress_bar import show_progress_bar


class ExcelExporter:
    '''
    ExcelExporter is a class that exports the pixel data to an excel file
    '''
    LINE_WIDTH = 4800
    LINE_COLOR = "0000FF"
    MAX_SHEET_NAME_LENGTH = 31

    def __init__(self, gray_img=None):
        self.gray_img = gray_img
    
    def export_to_excel(self, regions_data, filename):
        """Export pixel data to Excel file with separate sheets for each region"""
        try:
            # Add progress tracking for large exports
            total_regions = len(regions_data)
            progress_dialog = None
            
            if total_regions > 1:
                progress_dialog = wx.ProgressDialog(
                    "Exporting to Excel",
                    "Processing regions...",
                    maximum=100,
                    style=wx.PD_AUTO_HIDE | wx.PD_APP_MODAL
                )
            
            with pd.ExcelWriter(filename, engine='openpyxl') as writer:
                for idx, region_data in enumerate(regions_data):
                    region_id = region_data['region_id']
                    region_name = region_data.get('name', f'region_{region_id}')
                    region = region_data['region']
                    coords = region_data['coordinates']
                    max_points = region_data.get('max_points', 0)
                    
                    # Update progress
                    if progress_dialog:
                        progress = int((idx / total_regions) * 90)
                        progress_dialog.Update(progress, f"Processing {region_name}...")
                        wx.GetApp().Yield()
                    
                    # Process the region to get pixel data
                    df = self._process_region_data(region, coords, max_points)
                    
                    # Create sheet name (clean it for Excel compatibility)
                    sheet_name = region_name.replace('/', '_').replace('\\', '_').replace(':', '_').replace('?', '_').replace('*', '_').replace('[', '_').replace(']', '_')
                    if len(sheet_name) > self.MAX_SHEET_NAME_LENGTH:
                        sheet_name = f"R{region_id}_{sheet_name[:20]}"
                    
                    # Ensure sheet name is unique
                    original_name = sheet_name
                    counter = 1
                    while sheet_name in writer.sheets:
                        sheet_name = f"{original_name}_{counter}"
                        if len(sheet_name) > self.MAX_SHEET_NAME_LENGTH:
                            sheet_name = f"R{region_id}_{counter}"
                        counter += 1
                    
                    # Write to Excel
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
                    
                    # Get worksheet and add chart
                    worksheet = writer.sheets[sheet_name]
                    self._add_chart(worksheet, df, region_name, len(df))
                    
                    # Auto-adjust column widths
                    for column in worksheet.columns:
                        max_length = 0
                        column = [cell for cell in column]
                        for cell in column:
                            try:
                                if len(str(cell.value)) > max_length:
                                    max_length = len(str(cell.value))
                            except:
                                pass
                        adjusted_width = min(max_length + 2, 50)
                        worksheet.column_dimensions[column[0].column_letter].width = adjusted_width
                
                # Ensure all sheets are visible before saving
                if progress_dialog:
                    progress_dialog.Update(95, "Finalizing Excel file...")
                    wx.GetApp().Yield()
                
                # Make sure at least one sheet is visible
                workbook = writer.book
                for sheet in workbook.worksheets:
                    sheet.sheet_state = 'visible'
                
            # Close progress dialog
            if progress_dialog:
                progress_dialog.Update(100, "Export complete!")
                wx.GetApp().Yield()
                progress_dialog.Destroy()
                    
        except Exception as e:
            if progress_dialog:
                progress_dialog.Destroy()
            raise Exception(f"Failed to export to Excel: {str(e)}")

    def _process_region_data(self, region, coords, max_points):
        """Process region data to create DataFrame with pixel coordinates and grayscale values"""
        x, y, w, h = coords
        
        # Get all pixels in the region
        pixels = []
        for row in range(h):
            for col in range(w):
                pixel_x = x + col
                pixel_y = y + row
                grayscale_value = region[row, col]
                
                pixels.append({
                    'Index': len(pixels) + 1,
                    'X': pixel_x,
                    'Y': pixel_y,
                    'Grayscale': int(grayscale_value)
                })
        
        # Apply max_points limit if specified
        if max_points > 0 and len(pixels) > max_points:
            # Sample pixels uniformly
            indices = np.linspace(0, len(pixels) - 1, max_points, dtype=int)
            pixels = [pixels[i] for i in indices]
            # Update indices
            for i, pixel in enumerate(pixels):
                pixel['Index'] = i + 1
        
        return pd.DataFrame(pixels)


    def _process_rectangle(self, start, end, max_points):
        x1, y1 = start
        x2, y2 = end
        x1, x2 = sorted([max(0, min(x, self.gray_img.shape[1] - 1)) for x in (x1, x2)])
        y1, y2 = sorted([max(0, min(y, self.gray_img.shape[0] - 1)) for y in (y1, y2)])
        
        rect_pixels = self.gray_img[y1:y2+1, x1:x2+1].flatten()
        total_points = len(rect_pixels)

        if total_points > max_points and max_points != 0:
            indices = np.linspace(0, total_points - 1, max_points, dtype=int)
            rect_pixels = rect_pixels[indices]
        else:
            indices = np.arange(total_points)
            max_points = total_points

        x_coords = x1 + (indices % (x2 - x1 + 1))
        y_coords = self.gray_img.shape[0] - (y1 + (indices // (x2 - x1 + 1))) - 1

        df = pd.DataFrame({
            "Index": np.arange(1, len(indices) + 1),
            "Grayscale": rect_pixels,
            "X": x_coords,
            "Y": y_coords
        })

        return df, max_points

    def _add_chart(self, worksheet, df, name, max_points):
        chart = LineChart()
        chart.title = f"Grayscale Value vs Index - {name}"
        chart.style = 13
        chart.x_axis.title = "Index"
        chart.y_axis.title = "Grayscale Value"

        data = Reference(worksheet, min_col=2, min_row=1, max_col=2, max_row=len(df) + 1)
        categories = Reference(worksheet, min_col=1, min_row=2, max_row=len(df) + 1)

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.x_axis.scaling.max = max_points

        s = chart.series[0]
        s.graphicalProperties.line.solidFill = self.LINE_COLOR
        s.graphicalProperties.line.width = self.LINE_WIDTH

        worksheet.add_chart(chart, "F2")

    def _export_data(self, filename, rectangles, single_rectangle=False):
        if self.gray_img is None:
            messagebox.showerror("Error", "Please select an image first.")
            return

        if not filename:
            return

        def export_task():
            try:
                with pd.ExcelWriter(filename, engine="openpyxl") as writer:
                    rect_data = [rectangles] if single_rectangle else rectangles

                    for start, end, name, max_points in rect_data:
                        print(start, end, name, max_points)
                        df, max_points = self._process_rectangle(start, end, max_points)
                        sheet_name = name[:self.MAX_SHEET_NAME_LENGTH]  # Limit sheet name to 31 characters
                        df.to_excel(writer, sheet_name=sheet_name, index=False)
                        self._add_chart(writer.sheets[sheet_name], df, name, max_points)

                return True
            except Exception as e:
                print(f"Error exporting data: {str(e)}")
                return False

        result = show_progress_bar("Exporting Data......", export_task)

        if result:
            messagebox.showinfo("Success", "Data exported and charts added successfully!")
        else:
            messagebox.showerror("Error", "Failed to export data. Check console for details.")

    def export_all_data_to_excel(self, rectangles):
        filename = self._get_filename("data")
        self._export_data(filename, rectangles)

    def export_single_rectangle_data(self, start, end, name, max_points):
        filename = self._get_filename(f"{name}_data")
        self._export_data(filename, (start, end, name, max_points), single_rectangle=True)
